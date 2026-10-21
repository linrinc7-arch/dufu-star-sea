# -*- coding: utf-8 -*-
"""06_make_docs.py — 把两份说明转成 DOCX（宋体、小四、1.5 倍行距）"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(BASE, 'docs')

def set_font(run, name='宋体', size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)

def build(md_path, out_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0); sec.right_margin = Cm(3.0)

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(16)
            set_font(p.add_run(line[2:].strip()), '黑体', 18, True)
        elif line.startswith('## '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            set_font(p.add_run(line[3:].strip()), '黑体', 13, True)
        else:
            txt = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.first_line_indent = Pt(24)
            set_font(p.add_run(txt), '宋体', 12)
    doc.save(out_path)
    print('docx ->', os.path.basename(out_path))

build(os.path.join(DOCS,'技术方法说明.md'), os.path.join(DOCS,'技术方法说明.docx'))
build(os.path.join(DOCS,'经典文本关联说明.md'), os.path.join(DOCS,'经典文本关联说明.docx'))
